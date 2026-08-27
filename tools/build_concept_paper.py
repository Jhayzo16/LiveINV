from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Mabansag\OJT\InventorySystem")
OUT = ROOT / "deliverables" / "TGMCI_Mobile_Visual_Inventory_Tracking_System_Concept_Paper.docx"
LOGO = Path(r"C:\Users\Jhayzo\Downloads\Logo.png")

BURGUNDY = "8B1D24"
GREEN = "1B6C24"
DARK = "22262B"
MUTED = "66707A"
LIGHT = "F3F5F6"
PALE_GREEN = "EFF6F0"
PALE_RED = "F8EEEE"
WHITE = "FFFFFF"
FONT = "Arial"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=None, color=DARK, bold=None, italic=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_border(cell, color="D8DDE1", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(DARK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.33

    for name, size, color, before, after in (
        ("Heading 1", 16, BURGUNDY, 18, 10),
        ("Heading 2", 13, GREEN, 12, 6),
        ("Heading 3", 12, DARK, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        body = p.add_run(text[len(bold_lead):])
        set_run_font(body)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, size=10.5, bold=True, color=DARK)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p


def add_callout(doc, label, text, fill=PALE_GREEN, accent=GREEN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent=120)
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_border(cell, color=accent, size="8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    set_run_font(r, size=9, color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r = p2.add_run(text)
    set_run_font(r, size=10.5, color=DARK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GREEN)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    # Running header/footer. The first page uses the same restrained document furniture.
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("TGMCI IT DEPARTMENT  |  CONCEPT PAPER")
    set_run_font(hr, size=7.5, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])

    # Cover - editorial cover pattern with branded overrides.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    p.add_run().add_picture(str(LOGO), width=Inches(5.2))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("CONCEPT PAPER")
    set_run_font(r, size=10, color=GREEN, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("TGMCI Mobile Visual\nInventory Tracking System")
    set_run_font(r, size=28, color=BURGUNDY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("A React Native companion to the web-based hospital inventory platform")
    set_run_font(r, size=12.5, color=MUTED, italic=True)

    add_callout(
        doc,
        "Project Objective",
        "To design and develop a secure, cross-platform mobile inventory application that gives authorized hospital staff real-time, room-level visibility of assets and their assignment, condition, QR identity, maintenance, and network information.",
        fill=LIGHT,
        accent=BURGUNDY,
    )

    meta = doc.add_table(rows=4, cols=2)
    set_table_geometry(meta, [1900, 7460], indent=120)
    metadata = [
        ("Prepared by", "[Student Names]"),
        ("Course / Subject", "[Subject]"),
        ("Instructor", "[Instructor]"),
        ("Date", "[Date]"),
    ]
    for row, (label, value) in zip(meta.rows, metadata):
        shade(row.cells[0], PALE_GREEN)
        for cell in row.cells:
            set_cell_border(cell, color="DDE3E0", size="4")
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(label), size=9.5, color=GREEN, bold=True)
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), size=10, color=DARK)

    doc.add_page_break()

    add_heading(doc, "Concept Overview", 1)
    add_body(doc, "The proposed TGMCI Mobile Visual Inventory Tracking System extends the hospital's existing web-based inventory frontend into a mobile workflow for Android and, later, iOS. The mobile application is intended for authorized IT and operational personnel conducting equipment verification, assignment, transfer, and maintenance activities throughout the facility. It preserves the visual floor-to-room-to-asset concept while adding mobile camera access, QR scanning, protected offline work, and field-ready interaction.")
    add_callout(doc, "Scope Boundary", "The first mobile release is an IT inventory tool. Patient information, medical records, clinical decision support, and patient-care workflows are outside the project scope.", fill=PALE_RED, accent=BURGUNDY)

    add_heading(doc, "1. Problem", 1)
    add_body(doc, "Hospital IT inventory is commonly recorded through spreadsheets, paper forms, and disconnected files. These methods make it difficult to determine an asset's current floor, room, department, category, condition, custodian, QR identity, maintenance history, and network details such as hostname, IPv4 address, and MAC address. When several people update separate records, the institution can accumulate duplicate entries, inconsistent labels, and incomplete transfer histories.")
    add_body(doc, "The lack of room-level visibility can delay audits and make equipment difficult to locate. Unrecorded transfers can also cause IP address conflicts, inaccurate maintenance planning, and weak accountability for hospital-owned devices. A desktop-only workflow further limits staff during physical rounds because they must take notes and update the inventory later, increasing the chance of encoding errors and delayed status changes.")
    add_body(doc, "The problem is therefore not only the absence of a digital list, but the absence of one reliable, location-aware workflow that connects the physical asset, its QR identity, its assigned room, its operational condition, and its network record. The proposed project addresses equipment and technology infrastructure only; it does not process patient clinical records.")

    add_heading(doc, "2. Solution", 1)
    add_body(doc, "The proposed solution is a mobile-first companion to the TGMCI Visual Inventory Tracking System. Authorized users will be able to navigate from the hospital topology to a floor, select a room, inspect assigned items, scan or generate an asset QR code, and update inventory details at the point of inspection. Manual registration remains available for items without a readable label.")
    add_body(doc, "The mobile application and web administration interface will share a secure REST API and centralized database. This creates one source of truth while allowing each interface to serve a different work context: the web platform for broad administration, reporting, and map management; the mobile application for scanning, verification, assignment, and maintenance rounds.")
    add_callout(doc, "Recommended Architecture", "Use Clean Architecture implemented as feature-oriented layers: Presentation, Application / Use Cases, Domain, Data, and Infrastructure. This structure supports SOLID principles and keeps QR scanning, map rendering, authentication, storage, synchronization, and API access independently replaceable and testable.")

    add_heading(doc, "Architecture Principles", 2)
    principles = [
        ("Single Responsibility", "Each screen, use case, repository, and device service has one clear purpose."),
        ("Open / Closed", "New categories, status types, map renderers, or QR formats can be added without rewriting core inventory rules."),
        ("Liskov Substitution", "Mock, local, and remote repository implementations follow the same contracts."),
        ("Interface Segregation", "Camera, QR, authentication, map, network, and synchronization capabilities expose focused interfaces."),
        ("Dependency Inversion", "Business rules depend on abstractions rather than directly on Expo, storage, or HTTP libraries."),
    ]
    for label, detail in principles:
        add_bullet(doc, f"{label}. {detail}", bold_lead=f"{label}. ")

    add_heading(doc, "Connectivity Approach", 2)
    add_body(doc, "The app should synchronize normally when connected to the hospital network. When connectivity is temporarily unavailable, approved low-risk changes can be stored in an encrypted local queue with timestamps and user identifiers, then synchronized after reconnection. Conflict rules and server validation remain authoritative.")

    add_heading(doc, "3. Beneficiaries", 1)
    beneficiaries = [
        ("IT inventory personnel", "Faster rounds, camera-based QR verification, room-level lookup, category management, and a clearer IP/MAC registry."),
        ("Hospital management and administrators", "More reliable totals, condition summaries, accountability, utilization visibility, and audit-ready reports."),
        ("Clinical and administrative departments", "Faster transfers, clearer ownership of assigned equipment, and easier reporting of damaged or missing items."),
        ("Maintenance technicians", "Accessible status, notes, timestamps, warranty information, and previous maintenance activity before service begins."),
        ("Auditors and compliance officers", "Traceable assignment, transfer, status, and modification histories supported by authenticated audit logs."),
        ("Student developers and researchers", "A practical cross-platform project involving mobile development, system architecture, QR workflows, mapping, security, and synchronization."),
    ]
    for label, detail in beneficiaries:
        add_heading(doc, label, 2)
        add_body(doc, detail)
    add_callout(doc, "Access Model", "All capabilities are role-based. Users should only see or change information required by their assigned responsibilities, following least-privilege principles.", fill=LIGHT, accent=GREEN)

    add_heading(doc, "4. Features", 1)
    features = [
        ("Secure access", "Login, role-based authorization, session protection, and account activity logging."),
        ("Inventory dashboard", "Totals for active, maintenance, broken, and inactive assets, with filters and recent activity."),
        ("Visual topology", "Interactive hospital navigation from floors to rooms and from rooms to assigned assets."),
        ("Floor mapping", "Clickable and zoomable floor plans with hover or touch summaries for every mapped room."),
        ("QR workflow", "Camera-based scanning plus QR label generation for each registered item."),
        ("Registration and assignment", "Manual asset entry and controlled assignment or transfer to a floor, room, department, or custodian."),
        ("Categorization", "Computers, printers, monitors, networking devices, medical-support equipment, furniture, and configurable additional categories."),
        ("Asset profile", "Asset tag, serial number, brand, model, warranty, purchase date, assignment date, condition, and notes."),
        ("Network registry", "Hostname, IPv4 address, MAC address, VLAN or location where applicable, with validation to reduce duplicate addressing."),
        ("History and maintenance", "Status changes, maintenance actions, transfer records, notes, timestamps, and responsible user."),
        ("Search and notifications", "Search, categories, filters, recently viewed items, due-maintenance reminders, and change notifications."),
        ("Offline work", "Protected local queue and cache with controlled synchronization after connectivity returns."),
        ("Usability and privacy", "Accessible controls, clear camera-permission guidance, protected local data, and minimal device permissions."),
    ]
    for label, detail in features:
        add_bullet(doc, f"{label}: {detail}", bold_lead=f"{label}: ")

    add_heading(doc, "Primary Mobile Workflow", 2)
    flow = doc.add_table(rows=1, cols=4)
    set_table_geometry(flow, [2340, 2340, 2340, 2340], indent=120)
    for idx, (title, detail) in enumerate([
        ("1  Locate", "Choose floor and room from the visual map."),
        ("2  Identify", "Scan QR or search by asset tag."),
        ("3  Verify", "Check category, status, assignment, and network data."),
        ("4  Update", "Save assignment or maintenance changes and synchronize."),
    ]):
        cell = flow.cell(0, idx)
        shade(cell, PALE_GREEN if idx % 2 == 0 else LIGHT)
        set_cell_border(cell, color="D8E1DA", size="4")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(title), size=10, color=GREEN, bold=True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(detail), size=8.8, color=DARK)

    add_heading(doc, "5. Tech Stack", 1)
    tech = doc.add_table(rows=1, cols=3)
    set_table_geometry(tech, [2000, 3000, 4360], indent=120)
    headers = ["Layer", "Recommended Technology", "Purpose"]
    for idx, value in enumerate(headers):
        cell = tech.rows[0].cells[idx]
        shade(cell, BURGUNDY)
        set_cell_border(cell, color=BURGUNDY, size="6")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), size=9.5, color=WHITE, bold=True)
    set_repeat_table_header(tech.rows[0])

    rows = [
        ("Mobile client", "React Native + TypeScript", "Cross-platform Android/iOS interface with type-safe application code."),
        ("Framework", "Expo", "Development builds, native modules, camera access, updates, and delivery tooling."),
        ("Navigation", "Expo Router or React Navigation", "Typed screen flow for dashboard, floors, rooms, assets, scanner, and settings."),
        ("QR scanning", "expo-camera", "Camera preview and QR/barcode detection on supported devices."),
        ("Maps", "react-native-svg or React Native Skia", "Interactive topology, SVG floor plans, touch targets, zoom, and animation."),
        ("Client state", "Zustand or Redux Toolkit", "Authentication, selected location, filters, and draft changes."),
        ("Server state", "TanStack Query", "API caching, mutation state, retry, and synchronization coordination."),
        ("Offline storage", "Expo SQLite", "Local cache and queued changes during temporary network loss."),
        ("Secret storage", "Expo SecureStore", "Protected storage for session credentials and small sensitive values."),
        ("Backend", "Node.js + TypeScript; NestJS or Express", "Shared REST API for both web and mobile clients."),
        ("Database", "PostgreSQL + Prisma ORM", "Centralized assets, locations, assignments, users, maintenance, and audit records."),
        ("Authentication", "JWT/session + RBAC", "Authenticated sessions and least-privilege role enforcement."),
        ("Testing", "Jest, React Native Testing Library, Maestro or Detox", "Unit, component, workflow, and device-level validation."),
        ("Quality", "ESLint, Prettier, Git and GitHub", "Consistent code, collaboration, reviews, and version history."),
        ("Build and release", "Expo Application Services; Android Studio; later Xcode", "Development builds, signing, testing, and store-ready packages."),
    ]
    for row_index, values in enumerate(rows):
        row = tech.add_row()
        if row_index % 2:
            for cell in row.cells:
                shade(cell, "F8F9FA")
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            set_run_font(p.add_run(value), size=8.5, color=DARK, bold=(idx == 0))

    add_heading(doc, "Security and Privacy Requirements", 2)
    security = [
        "Use TLS for all API traffic and validate every request on the server.",
        "Apply least privilege through role-based authorization and maintain authenticated audit logs.",
        "Store session credentials through platform-protected secure storage rather than plain local files.",
        "Minimize camera and device permissions; explain why each permission is requested.",
        "Protect offline data, define retention rules, and provide backup and recovery procedures.",
        "Use OWASP MASVS as the mobile security baseline for storage, authentication, network communication, platform use, code quality, resilience, and privacy.",
        "Conduct a Philippine Data Privacy Act compliance review if personal or sensitive personal information is introduced in any future scope.",
    ]
    for item in security:
        add_bullet(doc, item)

    add_heading(doc, "6. References", 1)
    references = [
        ("Meta Platforms, Inc. (2026). React Native documentation: Set Up Your Environment. ", "https://reactnative.dev/docs/set-up-your-environment"),
        ("Expo. (2026a). Introduction to Expo Router. ", "https://docs.expo.dev/router/introduction/"),
        ("Expo. (2026b). Camera. ", "https://docs.expo.dev/versions/latest/sdk/camera/"),
        ("React Navigation. (2026). Getting started. ", "https://reactnavigation.org/docs/getting-started/"),
        ("OWASP Foundation. (2026). Mobile Application Security Verification Standard (MASVS). ", "https://mas.owasp.org/MASVS/"),
        ("National Privacy Commission. (2012). Republic Act No. 10173 - Data Privacy Act of 2012. ", "https://privacy.gov.ph/data-privacy-act/"),
    ]
    for lead, url in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.2
        set_run_font(p.add_run(lead), size=10)
        add_hyperlink(p, url, url)

    add_heading(doc, "Expected Outcome", 1)
    add_body(doc, "The expected result is a validated Android-first prototype that reduces manual lookup time, improves location accuracy and accountability, supports room-level audits, and establishes a reusable shared architecture for web and mobile deployment. The project will demonstrate that a visually mapped, QR-enabled inventory workflow can help authorized hospital personnel verify and update equipment information where the assets are physically located.")
    add_callout(doc, "Success Definition", "A secure and usable prototype in which an authorized user can locate a room, scan or find an asset, verify its assignment and condition, update permitted information, and synchronize the change to the shared inventory service.", fill=PALE_GREEN, accent=GREEN)

    # Set document core properties without personal metadata.
    doc.core_properties.title = "TGMCI Mobile Visual Inventory Tracking System - Concept Paper"
    doc.core_properties.subject = "React Native hospital IT inventory system concept paper"
    doc.core_properties.author = "TGMCI IT Department Project Team"
    doc.core_properties.keywords = "inventory, React Native, QR code, hospital, floor map, mobile app"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
