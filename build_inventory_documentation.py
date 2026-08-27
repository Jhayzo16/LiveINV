from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = r"C:\Mabansag\OJT\InventorySystem\Hospital_Visual_Inventory_Tracking_System_Final_Documentation.docx"

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT = "D9EAF7"
PALE = "F4F8FC"
GRAY = "5B6573"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.75)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Plus Jakarta Sans'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Plus Jakarta Sans')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Plus Jakarta Sans')
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Plus Jakarta Sans')
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size, color, before, after in [('Title', 25, NAVY, 0, 10), ('Subtitle', 12, GRAY, 0, 18), ('Heading 1', 16, NAVY, 15, 7), ('Heading 2', 12, BLUE, 10, 5), ('Heading 3', 10.5, NAVY, 7, 3)]:
    s = styles[name]
    s.font.name = 'Montserrat' if name in ('Title','Heading 1','Heading 2','Heading 3') else 'Plus Jakarta Sans'
    s._element.rPr.rFonts.set(qn('w:ascii'), s.font.name)
    s._element.rPr.rFonts.set(qn('w:hAnsi'), s.font.name)
    s._element.rPr.rFonts.set(qn('w:eastAsia'), s.font.name)
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = name != 'Subtitle'
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None: node = OxmlElement(f'w:{side}'); mar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def set_width(cell, width):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None: tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width)); tcW.set(qn('w:type'), 'dxa')

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c, NAVY); set_cell_margin(c)
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
        if widths: set_width(c,widths[i])
    trPr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    trPr.append(tbl_header)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; c.text=str(val); set_cell_margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==0: shade(c, PALE)
            if widths: set_width(c,widths[i])
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(1)
                for r in p.runs: r.font.size=Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def p(text='', style=None, bold_prefix=None):
    par=doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        par.add_run(bold_prefix).bold=True; par.add_run(text[len(bold_prefix):])
    else: par.add_run(text)
    return par

def bullets(items):
    for x in items: p(x, 'List Bullet')

def numbered(items):
    for x in items: p(x, 'List Number')

def callout(label, text):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False
    c=t.cell(0,0); shade(c,LIGHT); set_cell_margin(c,130,160,130,160); set_width(c,9000)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    trPr.append(tbl_header)
    para=c.paragraphs[0]; para.paragraph_format.space_after=Pt(0)
    r=para.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
    para.add_run(text)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def h(text, level=1): doc.add_heading(text, level=level)
def page(): doc.add_page_break()

# Header/footer
header=sec.header.paragraphs[0]
header.text='HOSPITAL VISUAL INVENTORY TRACKING SYSTEM  |  FRONT-END BLUEPRINT'
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Internal project documentation  |  August 2026  |  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
for r in footer.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

# Cover
p('HOSPITAL VISUAL INVENTORY', 'Title')
p('TRACKING SYSTEM', 'Title')
p('Final Documentation and Front-End Implementation Blueprint', 'Subtitle')
p('A visual, room-aware asset inventory system for tracking IT equipment and hospital assets by building, floor, department, room, and item.', bold_prefix=None)
doc.add_paragraph()
table(['Document control','Value'],[
 ['System name','Hospital Visual Inventory Tracking System'],
 ['Release focus','Web-system UI expansion and architecture baseline'],
 ['Primary users','System administrators, IT / inventory staff, department viewers'],
 ['Prepared for','Hospital inventory modernization project'],
 ['Version','1.1 - Architecture and UI examples update'],
 ['Date','25 August 2026'],
], [2800,6200])
callout('DECISION', 'Use a modular monolith with layered Clean Architecture. Organize React by feature and the NestJS API by business module; keep controllers, use cases, domain rules and repositories separate so the first deployment remains simple without locking the project into MVC or microservices.')
h('Document purpose',1)
p('This document finalizes the functional scope, front-end architecture, data contracts, QR workflows, room and department assignment rules, PC IP-address handling, and recommended project structure. It is written so the first UI can be built before the backend and database are available.')
page()

h('1. Overview and objectives')
p('The system makes hospital inventory easy to locate visually. Staff navigate from hospital to floor, then department or room, then to a specific item. Each item retains its category, status, assigned location, details, and movement history.')
h('Objectives',2)
bullets([
 'Provide a single visual inventory workspace for IT and non-IT assets.',
 'Assign each asset to an exact floor, department, and room using manual selection or QR scanning.',
 'Create a unique QR code for every item and open its details with a camera scan.',
 'Record PC network details, including assigned IP address, hostname, MAC address, and network status.',
 'Organize assets by category, subcategory, condition, and ownership.',
 'Keep the floor plan as a visual layer; inventory and location data remain the source of truth.',
 'Deliver the initial version as a responsive front-end prototype backed by controlled mock data.'
])
h('Initial scope',2)
table(['Included in phase 1','Deferred until backend phase'],[
 ['Interactive building / floor selector; SVG floor map and clickable rooms','Authentication, authorization, persistent database and real audit storage'],
 ['Asset listing, search, filters, details, categories, manual assignment and QR UI','Server-side QR validation, live camera permissions policy and API integration'],
 ['PC network-detail form and IP-address registry view','Automatic hardware discovery, network scanning, GPS, patient records and consumables'],
 ['QR preview, print/download actions and scan simulation','Permanent transfer history, maintenance workflows and reporting exports'],
], [4500,4500])
h('Roles',2)
table(['Role','Front-end permissions'],[
 ['System Administrator','Configure floors, rooms, departments, categories, QR labels and all inventory screens.'],
 ['IT / Inventory Staff','Add and edit items, scan or generate QR codes, assign or move assets, maintain PC details.'],
 ['Department Viewer','Search and view approved inventory, rooms and item detail; no edit actions.'],
], [2500,6500])
page()

h('2. Information model and assignment rules')
h('Location hierarchy',2)
p('Every assigned item follows this hierarchy: Hospital -> Floor -> Department -> Room -> Asset. Department is required when a room belongs to a department; shared rooms may be assigned to an operational area such as IT, Facilities, or Shared Services.')
table(['Entity','Required fields for the prototype'],[
 ['Floor','id, number, name, map asset, status'],
 ['Department','id, name, code, floorIds (or rooms managed), contact person'],
 ['Room','id, floorId, departmentId, code, name, type, SVG region id'],
 ['Category','id, name, parentCategoryId, icon, active flag'],
 ['Asset','id, assetTag, displayName, categoryId, status, serialNumber, brand, model'],
 ['Assignment','assetId, floorId, departmentId, roomId, assignedAt, assignedBy, method'],
 ['PC network profile','assetId, hostname, IPv4 address, MAC address, VLAN / network, DHCP or static'],
 ['QR identity','assetId, qrToken, labelText, generatedAt, label status'],
], [2400,6600])
callout('DATA RULE', 'Store the current room assignment separately from the display map. An item can be drawn on a map only after it has a valid floor and room. A later backend transfer will create history; the front end should already model a pending assignment event.')
h('Category model',2)
table(['Top-level category','Example subcategories'],[
 ['Computing','Desktop PC, laptop, monitor, keyboard, mouse, UPS'],
 ['Printing and scanning','Network printer, label printer, scanner, photocopier'],
 ['Networking','Switch, router, access point, firewall, patch panel'],
 ['Clinical / facility equipment','Wheelchair, infusion equipment, refrigerator, generator component'],
 ['Office and furniture','Desk, cabinet, filing system, telephone'],
 ['Software / licenses','Operating system, office suite, endpoint protection'],
], [3200,5800])
page()

h('3. Core front-end user flows')
h('Manual assignment',2)
numbered([
 'Staff opens an item from search, a category list, or its details page.',
 'Staff selects Assign / Move Item and chooses Floor, Department, and Room in that order.',
 'The UI filters departments by floor and rooms by the selected department and floor.',
 'The user reviews the target location, condition, custodian and reason, then confirms.',
 'The asset detail, room count and floor-map marker update immediately in the prototype.'
])
h('QR scan and QR creation',2)
table(['Workflow','Front-end behavior'],[
 ['Create QR per item','Generate a QR value that contains an opaque item token or a route such as /assets/{assetId}; show label preview, printable text and download action. Do not embed sensitive PC credentials in QR data.'],
 ['Scan to view details','Open a camera-scanner screen (or manual-code fallback). On a valid scan, route to the asset detail page and show current category, location, status and PC/IP information when authorized.'],
 ['Scan to assign','Scan the item first, then select the target floor / department / room. The confirmation screen clearly states the old and new locations and records method = QR.'],
 ['Manual code entry','Support an asset-tag or QR-token input for devices without a camera or unreadable labels.'],
], [2500,6500])
callout('QR SAFETY', 'The QR label should identify the asset, not expose passwords, Wi-Fi keys, patient information, or full network configuration. Access to sensitive fields remains role-based when the detail screen opens.')
h('Asset detail screen',2)
bullets([
 'Identity: asset tag, QR label, category, brand, model, serial number and image.',
 'Location: building, floor, department, room, assigned date and assignment method.',
 'Status: active, maintenance, broken, inactive, under transfer or retired.',
 'PC-only details: hostname, IPv4 address, MAC address, network/VLAN, operating system, CPU, RAM and storage.',
 'Activity panel: mock assignment and status timeline in phase 1; durable audit history in the backend phase.'
])
page()

h('4. PC IP-address registry')
p('IP addresses are configuration data for computer and network-capable assets. The system should show the actual address only to authorized IT users. Because no real hospital addressing plan was supplied, the entries below are a safe planning template, not production addresses.')
table(['Field','Rule / example'],[
 ['Asset tag','Required and unique, e.g., IT-PC-0034'],
 ['Hostname','Required for managed PCs, e.g., HOSP-MEDREC-034'],
 ['IPv4 address','Validate IPv4 format and uniqueness within a network; use a supplied real address only after IT confirmation.'],
 ['Address method','DHCP reservation, static, or not applicable.'],
 ['MAC address','Optional in phase 1; validate hexadecimal format when present.'],
 ['Network / VLAN','Human-readable network name and optional VLAN number.'],
 ['Last verified','Date and technician; clearly mark mock/demo values in the prototype.'],
], [2500,6500])
h('Example registry (placeholder values only)',2)
table(['Asset tag','Hostname','IP address','Method','Location'],[
 ['IT-PC-0034','HOSP-MEDREC-034','10.20.3.34','DHCP reservation','Floor 3 / Medical Records'],
 ['IT-PC-0102','HOSP-ER-102','10.20.1.102','Static','Floor 1 / Emergency Room'],
 ['IT-PRN-0018','HOSP-FIN-PRN-18','10.20.4.18','Static','Floor 4 / Finance'],
], [1600,2200,1600,1600,2000])
callout('IMPLEMENTATION NOTE', 'Use reserved private ranges in demo data only. The final list of actual PC IP addresses must be provided or approved by hospital IT, then loaded through an admin import or secure API.')
h('Validation rules',2)
bullets([
 'An IPv4 address cannot appear twice on the same network unless the prior record is retired.',
 'A PC can exist without an IP address when it is offline, newly received or not network-capable.',
 'Network data is hidden from Viewer users and shown in read-only form unless the staff role can edit it.',
 'The IP registry filters by floor, department, room, category, status and verification date.'
])
page()

h('5. Recommended architecture')
h('Recommendation: modular monolith with layered Clean Architecture',2)
p('Use one deployable web application and one deployable API, divided internally into business modules such as assets, locations, assignments, QR, network registry, maintenance, reports and users. Within each module, dependencies move inward from controllers and UI adapters to application use cases, then to domain rules. Repository interfaces shield the domain from Prisma, PostgreSQL, cameras, QR libraries and other infrastructure. This is more suitable than strict MVC because React is component-driven and the system contains business workflows that should not live in controllers or views.')
table(['Layer','Responsibility','Examples'],[
 ['Presentation / delivery','React pages and components; NestJS controllers and DTOs. No authoritative business rules.','Dashboard, asset table, REST controllers'],
 ['Application','Use cases, orchestration, authorization checks and transaction boundaries.','AssignAsset, ResolveQrToken, CloseWorkOrder'],
 ['Domain','Entities, value objects, invariants and repository interfaces.','Asset, Location, Assignment, NetworkProfile'],
 ['Infrastructure','Concrete adapters for persistence and external capabilities.','Prisma repositories, QR generator, camera/API adapters'],
 ['Shared platform','Cross-cutting contracts, logging, errors, design tokens and security utilities.','API contracts, audit context, UI primitives'],
], [1900,3900,3200])
h('Why this is the best fit',2)
bullets([
 'A modular monolith is simpler to deploy, secure, back up and monitor than microservices for the project’s current team and scale.',
 'Business modules remain independently testable and can be separated later only when real scaling or ownership needs appear.',
 'The web UI can use mock repositories now and switch to REST adapters without changing page-level workflows.',
 'PostgreSQL fits strongly related asset, location, assignment, maintenance and audit records; Prisma provides typed persistence at the infrastructure edge.',
 'MVC remains useful at the API boundary, but controllers stay thin and delegate to application use cases.'
])
h('Future deployment shape',2)
p('React + TypeScript + Vite web client -> versioned REST API -> NestJS modular monolith -> PostgreSQL through Prisma. Use JWT-based authentication with role-based access control, server-side validation, audit logging and TLS. SVG floor plans remain front-end assets, while floors, rooms, assets, assignments, work orders and history become persisted relational records. Split a module into a separate service only after measured load, release cadence or ownership justifies the operational cost.')
page()

h('6. SOLID design rules')
table(['Principle','Application in this system'],[
 ['Single Responsibility','FloorMap renders map interaction; AssignmentForm collects location; assignment use case validates and submits; repository stores data.'],
 ['Open / Closed','Add a new asset category or QR provider through configuration / adapters rather than editing every screen.'],
 ['Liskov Substitution','Mock repositories and API repositories conform to the same interfaces and can be exchanged safely.'],
 ['Interface Segregation','Use focused interfaces such as AssetReader, AssetWriter and QrCodeGenerator instead of one large all-purpose service.'],
 ['Dependency Inversion','Features depend on domain interfaces; concrete HTTP, camera and QR libraries are injected at the infrastructure edge.'],
], [2100,6900])
h('Front-end technical choices',2)
table(['Concern','Recommended choice'],[
 ['Framework','React 19 + TypeScript + Vite'],
 ['API architecture','NestJS modular monolith with thin controllers and application use cases'],
 ['Persistence','PostgreSQL through Prisma repositories and explicit database migrations'],
 ['Routing','React Router with protected-route placeholders'],
 ['Server state','TanStack Query once API is available; local mock query adapter initially'],
 ['Client UI state','Feature-local state first; Zustand only for truly shared UI state such as selected map context'],
 ['Forms and validation','React Hook Form + Zod schemas'],
 ['Map','SVG floor-plan components with room IDs and controlled marker overlays'],
 ['QR','A QR generation library; browser camera scanner behind a QrScanner interface'],
 ['Typography','Plus Jakarta Sans for body text, labels, fields and buttons; Montserrat for page and section titles'],
 ['Styling','Feature-level CSS with shared hospital design tokens; light operational workspace with maroon and green accents'],
 ['Testing','Vitest + React Testing Library; Playwright for key flows'],
], [2700,6300])
page()

h('7. Proposed project structure')
p('Organize the codebase as a small workspace with a web application, an API application and shared contracts. Within both applications, group code by business feature rather than by one global controllers, services or components folder.')
code = '''apps/
  web/src/
    app/ routes/ providers/
    pages/                 # route-level compositions
    features/
      assets/ assignments/ locations/ qr/
      network/ maintenance/ reports/ users/
    shared/ ui/ lib/ styles/
    infrastructure/ mock/ http/ camera/
  api/src/
    modules/
      assets/
        presentation/     # controllers and DTOs
        application/      # commands, queries, use cases
        domain/           # entities, rules, interfaces
        infrastructure/   # Prisma repository adapters
      locations/ assignments/ qr/
      network/ maintenance/ reports/ users/
    platform/ auth/ audit/ database/ errors/
packages/
  contracts/              # shared API schemas and enums
prisma/
  schema.prisma migrations/ seed/'''
for line in code.splitlines():
    par=p(line); par.paragraph_format.space_after=Pt(0); par.paragraph_format.line_spacing=1
    for r in par.runs: r.font.name='Cascadia Mono'; r.font.size=Pt(8.4)
h('Key contracts',2)
table(['Contract','Purpose'],[
 ['AssetRepository','Search, read, create/update and assign assets; mock implementation now, HTTP implementation later.'],
 ['LocationRepository','Read floors, departments and rooms; resolve valid room options for an assignment.'],
 ['QrCodeGenerator','Creates label data / image from a non-sensitive asset identity.'],
 ['QrScanner','Captures a code from camera or manual input and returns a normalized token.'],
 ['NetworkProfileValidator','Checks IP, MAC and hostname fields before saving a PC profile.'],
], [2600,6400])
page()

h('8. Screen inventory and UI acceptance criteria')
table(['Screen','Must support in phase 1'],[
 ['Dashboard','Counts by category, status and floor; recent mock activity; search entry point.'],
 ['Inventory list','Search by asset tag, name, serial number or QR code; filter category, floor, department, room and status.'],
 ['Floor map','Floor selector; SVG rooms with accessible labels; room tooltip and assigned-item count.'],
 ['Room / department panel','Room details, department name, item list, category breakdown and Assign Item action.'],
 ['Asset details','Identity, category, location, status, QR label and conditional PC network panel.'],
 ['Assign / move dialog','Manual hierarchy selection, optional QR-first flow, confirmation summary and instant mock update.'],
 ['QR label and scanner','Generate label preview, print/download action, camera permission states and manual-entry fallback.'],
 ['Network registry','IT-only list of PC and printer IP profiles with location and verification filters.'],
 ['Maintenance','Priority metrics, active work orders, preventive schedule, parts status and SLA visibility.'],
 ['Reports','Inventory, network, maintenance and movement reports with export history and basic visual summaries.'],
 ['Users and roles','User directory, account status and role-permission review based on least privilege.'],
], [2700,6300])
h('UI design standard and example-page direction',2)
bullets([
 'Use Plus Jakarta Sans for body copy, table content, labels, inputs and buttons. Use Montserrat for page titles, card titles and major section headings.',
 'Use a light operational workspace with white data surfaces, soft gray structure, hospital maroon for primary actions and hospital green for positive status.',
 'Keep the existing topology experience. Apply the shared shell, typography, spacing and status language to Dashboard, Assets, Assignments, QR Scanner, Network Registry, Maintenance, Reports and Users.',
 'The August 2026 example pages are design prototypes for review. They establish hierarchy and interaction patterns but do not yet imply persistent backend behavior.',
 'The supplied dark reference was used only to identify useful page categories and information density. The implemented visual composition is original.'
])
h('Acceptance checks',2)
bullets([
 'A user can add a mock asset, place it in a category and assign it to a valid floor / department / room.',
 'A scanned or manually entered QR token opens the matching asset and supports reassignment.',
 'Asset movement updates the detail page, inventory list and relevant room count without a browser refresh.',
 'A PC asset shows validated network fields; non-PC assets do not show irrelevant network inputs.',
 'Keyboard users can select floor-map rooms and receive the same room details as mouse users.',
 'The application has loading, empty, error and permission-denied states for every data-dependent screen.'
])

h('9. Data contracts for the front-end prototype')
table(['Type','Selected fields'],[
 ['Asset','id, assetTag, name, categoryId, status, serialNumber, brand, model, assignment, networkProfile?, qrToken'],
 ['Assignment','floorId, departmentId, roomId, assignedAt, assignedBy, method: manual | qr'],
 ['Room','id, floorId, departmentId, code, name, type, svgRegionId'],
 ['Category','id, name, parentId?, icon, active'],
 ['NetworkProfile','hostname, ipv4Address?, macAddress?, addressMethod, networkName?, verifiedAt?'],
 ['QrPayload','version, assetId or opaque token, issuedAt; never passwords or sensitive medical data'],
], [2100,6900])
h('Suggested API surface for later integration',2)
table(['Method','Endpoint','Purpose'],[
 ['GET','/api/floors','List floor summaries and maps'],
 ['GET','/api/floors/:floorId/rooms','Return room geometry and assignment counts'],
 ['GET','/api/assets','Search/filter inventory'],
 ['GET','/api/assets/:assetId','Read item details'],
 ['POST','/api/assets','Create an item'],
 ['PATCH','/api/assets/:assetId','Update metadata and category'],
 ['POST','/api/assets/:assetId/assignments','Assign/move item; create durable history'],
 ['POST','/api/assets/:assetId/qr','Create/rotate QR identity'],
 ['POST','/api/qr/resolve','Resolve a scanned token safely'],
 ['GET','/api/network-profiles','Authorized PC/IP registry search'],
], [1400,3600,4400])
callout('BACKEND BOUNDARY', 'The UI must not calculate authoritative permissions, room availability or history. It may validate usability rules, but the later API repeats and enforces all business-critical checks.')
page()

h('10. Implementation roadmap')
table(['Sprint / phase','Outcome'],[
 ['1. Foundation','Vite project, design tokens, routes, layout, mock data, domain types and repository interfaces.'],
 ['2. Inventory and locations','Asset list/detail, category management UI, floor/departments/rooms and manual assignment.'],
 ['3. Visual floor map','SVG map shell, room overlays, accessible selection, room inventory panel and markers.'],
 ['4. QR and network details','QR label creation, scanner mock/camera adapter, PC network profile and IP registry.'],
 ['5. Polish and readiness','Form validation, empty/error states, responsive QA, unit/component/E2E tests and API handoff contracts.'],
 ['6. Backend phase','Authentication, database migrations, API adapters, persistent QR, audit history and role enforcement.'],
], [2200,6800])
h('Testing priorities',2)
bullets([
 'Unit-test assignment validation, IP/MAC validation and QR-token normalization.',
 'Component-test filters, assignment confirmation and unauthorized network-profile visibility.',
 'End-to-end test manual assignment, QR lookup, QR reassignment and floor-map room selection.',
 'Use mock data with multiple floors, departments, categories, statuses and unassigned assets.'
])
h('Conclusion',2)
p('The recommended path is a modular monolith with layered Clean Architecture: a feature-organized React web client, a module-organized NestJS API and PostgreSQL persistence through Prisma. This gives the team a straightforward first deployment while keeping domain rules testable and infrastructure replaceable. The core rule remains simple: every asset has a category and a current location; QR makes retrieval and assignment fast; maintenance and audit events preserve accountability; and network data remains controlled, IT-only information.')

doc.core_properties.title = 'Hospital Visual Inventory Tracking System - Final Documentation'
doc.core_properties.subject = 'Front-end architecture, QR assignment, location and IP registry blueprint'
doc.core_properties.author = 'Hospital Inventory Project'
doc.save(OUT)
print(OUT)
